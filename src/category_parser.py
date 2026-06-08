from __future__ import annotations

import logging

log = logging.getLogger(__name__)

# Column order in the Tableau Category Performance export
CAT_COLS = ["main_cat", "sub_cat1", "sub_cat2", "sub_cat3", "sub_cat4"]


def parse_category_file(path: str) -> list[dict]:
    """
    Parse the Tableau Category Performance xlsx into flat rows.

    The export uses merged cells for parent categories (openpyxl returns None
    for the merged-away cells), so we forward-fill None values per column.
    Empty string ('') means the level genuinely does not apply (leaf is higher).

    Returns a list of dicts with: main_cat, sub_cat1..4, leaf_cat, level,
    gmv, gp, gp_pct.
    """
    import openpyxl

    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.worksheets[0]

    rows: list[dict] = []
    last = {c: "" for c in CAT_COLS}

    for r in range(2, ws.max_row + 1):  # row 1 is the header
        cells = [ws.cell(r, c).value for c in range(1, 9)]
        cat_vals = cells[0:5]
        gmv, gp, gp_pct = cells[5], cells[6], cells[7]

        # Skip fully blank rows
        if all(v is None or v == "" for v in cat_vals) and gmv is None:
            continue

        # Forward-fill None (merged cells); keep '' as a real empty level
        resolved = {}
        for i, col in enumerate(CAT_COLS):
            v = cat_vals[i]
            if v is None:
                v = last[col]
            else:
                last[col] = v
            resolved[col] = (v or "").strip()

        # leaf = deepest non-empty category level
        leaf = ""
        level = 0
        for i, col in enumerate(CAT_COLS):
            if resolved[col]:
                leaf = resolved[col]
                level = i + 1

        rows.append({
            **resolved,
            "leaf_cat": leaf,
            "level": level,
            "gmv": _num(gmv),
            "gp": _num(gp),
            "gp_pct": _num(gp_pct),
        })

    log.info("Parsed %d category rows from %s", len(rows), path)
    return rows


def _num(v) -> float:
    if v is None or v == "":
        return 0.0
    try:
        return float(v)
    except (ValueError, TypeError):
        return 0.0
